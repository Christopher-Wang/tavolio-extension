import csv
import random
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font


ROOT = Path(__file__).resolve().parent
SOURCE = Path(sys.argv[1])
OUTPUT = Path(sys.argv[2])
SEED = 20261002


def parse_arff(path):
    attributes = []
    rows = []
    in_data = False
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("%"):
            continue
        if line.lower() == "@data":
            in_data = True
            continue
        if not in_data:
            match = re.match(r"@attribute\s+('[^']+'|\S+)\s+(.+)", line, re.I)
            if match:
                attributes.append((match.group(1).strip("'"), match.group(2).strip()))
            continue
        values = next(csv.reader([line], quotechar="'", skipinitialspace=True))
        if len(values) != len(attributes):
            raise ValueError(f"Expected {len(attributes)} fields, found {len(values)}")
        record = []
        for value, (_name, kind) in zip(values, attributes):
            if value == "?":
                record.append(None)
            elif kind.upper() == "INTEGER":
                record.append(int(value))
            elif kind.upper() in {"REAL", "NUMERIC"}:
                record.append(float(value))
            else:
                record.append(value)
        rows.append(record)
    return [name for name, _kind in attributes], rows


headers, rows = parse_arff(ROOT / "data" / "maternal_health_risk.arff")
assert len(headers) == 7 and headers[-1] == "RiskLevel"
assert len(rows) == 1014
assert set(row[-1] for row in rows) == {"low risk", "mid risk", "high risk"}
original_labels = Counter(row[-1] for row in rows)

rng = random.Random(SEED)
blank_count = round(len(rows) * .05)
for index in rng.sample(range(len(rows)), blank_count):
    rows[index][-1] = None
rng.shuffle(rows)

wb = load_workbook(SOURCE)
assert "Marketing campaign" not in wb.sheetnames
assert "Maternal health risk" not in wb.sheetnames
ws = wb.create_sheet("Maternal health risk", wb.sheetnames.index("Sources & notes"))
for col, header in enumerate(headers, 1):
    cell = ws.cell(1, col, header)
    cell.font = Font(bold=True)
for row_number, values in enumerate(rows, 2):
    for col_number, value in enumerate(values, 1):
        ws.cell(row_number, col_number, value)

notes = wb["Sources & notes"]
method_row = next((r for r in range(2, notes.max_row + 1) if notes.cell(r, 1).value == "Method"), notes.max_row + 1)
notes.insert_rows(method_row)
note = [
    "Maternal health risk",
    "RiskLevel",
    1014,
    "https://www.openml.org/d/46600",
    f"OpenML task 363405. Blanked {blank_count} of 1,014 RiskLevel labels (5%) with seed {SEED}; shuffled all rows. Six source features unchanged. Task: https://www.openml.org/t/363405",
]
for col, value in enumerate(note, 1):
    notes.cell(method_row, col, value)
notes.auto_filter.ref = f"A1:E{notes.max_row}"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT)

saved = load_workbook(OUTPUT, read_only=True)
sheet = saved["Maternal health risk"]
iterator = sheet.iter_rows(values_only=True)
assert list(next(iterator)) == headers
saved_rows = list(iterator)
assert len(saved_rows) == 1014
saved_labels = Counter(row[-1] for row in saved_rows)
assert saved_labels[None] == blank_count
assert sum(saved_labels.values()) == 1014
assert len(saved.sheetnames) == 7
print("source labels:", dict(original_labels))
print("saved labels:", dict(saved_labels))
print("rows:", len(saved_rows), "features:", len(headers) - 1)
print("sheets:", saved.sheetnames)
print("output:", OUTPUT)
