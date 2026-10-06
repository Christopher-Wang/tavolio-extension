import csv
import random
import re
import sys
from collections import Counter
from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
SOURCE = Path(sys.argv[1])
OUTPUT = Path(sys.argv[2])
SEED = 20261002
BLANK_SHARE = 0.05


def parse_arff(path):
    attributes = []
    records = []
    in_data = False
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("%"):
            continue
        if line.lower() == "@data":
            in_data = True
            continue
        if not in_data:
            match = re.match(r"@attribute\s+('[^']+'|\S+)\s+(.+)", line, re.I)
            if match:
                attributes.append((match.group(1).strip("'"), match.group(2).strip()))
            continue
        values = next(csv.reader([line], quotechar="'", skipinitialspace=True))
        if len(values) != len(attributes):
            raise ValueError(f"Expected {len(attributes)} fields, found {len(values)}: {line[:100]}")
        record = []
        for value, (name, kind) in zip(values, attributes):
            value = value.strip()
            if value == "?":
                parsed = None
            elif name == "Dt_Customer":
                parsed = date.fromisoformat(value)
            elif kind.upper() == "INTEGER":
                parsed = int(value)
            elif kind.upper() in {"REAL", "NUMERIC"}:
                parsed = float(value)
            else:
                parsed = value
            record.append(parsed)
        records.append(record)
    return [name for name, _kind in attributes], records


headers, rows = parse_arff(ROOT / "data" / "marketing_campaign.arff")
assert len(headers) == 26, len(headers)
assert len(rows) == 2240, len(rows)
assert headers[-1] == "Response", headers[-1]
assert set(row[-1] for row in rows) == {"No", "Yes"}

original_labels = Counter(row[-1] for row in rows)
rng = random.Random(SEED)
blank_count = round(len(rows) * BLANK_SHARE)
for index in rng.sample(range(len(rows)), blank_count):
    rows[index][-1] = None
rng.shuffle(rows)

wb = load_workbook(SOURCE)
if "Marketing campaign" in wb:
    raise ValueError("Marketing campaign sheet already exists")
position = wb.sheetnames.index("Sources & notes")
ws = wb.create_sheet("Marketing campaign", position)
ws.sheet_view.showGridLines = False
ws.freeze_panes = "A2"
ws.sheet_properties.tabColor = "6B7280"

reference = wb["House prices"]
for col, name in enumerate(headers, 1):
    cell = ws.cell(1, col, name)
    cell._style = copy(reference.cell(1, 1)._style)
    cell.alignment = copy(reference.cell(1, 1).alignment)
    ws.column_dimensions[get_column_letter(col)].width = min(29, max(13, len(name) + 3))
ws.cell(1, len(headers))._style = copy(reference.cell(1, reference.max_column)._style)
ws.row_dimensions[1].height = 25

blank_fill = PatternFill("solid", fgColor="FFF2CC")
target_font = Font(name="Aptos", size=10, bold=True, color="14532D")
for row_number, values in enumerate(rows, 2):
    for col_number, value in enumerate(values, 1):
        cell = ws.cell(row_number, col_number, value)
        if isinstance(value, date):
            cell.number_format = "yyyy-mm-dd"
        if col_number == len(headers):
            cell.font = target_font
            if value is None:
                cell.fill = blank_fill
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

notes = wb["Sources & notes"]
method_row = next((r for r in range(2, notes.max_row + 1) if notes.cell(r, 1).value == "Method"), notes.max_row + 1)
notes.insert_rows(method_row)
notes_values = [
    "Marketing campaign",
    "Response",
    len(rows),
    "https://www.openml.org/d/46940",
    f"OpenML task 363647. Blanked {blank_count} of 2,240 Response labels (5%) with seed {SEED}; shuffled all rows. Source feature values and existing nulls retained. Task: https://www.openml.org/t/363647",
]
for col, value in enumerate(notes_values, 1):
    cell = notes.cell(method_row, col, value)
    cell._style = copy(notes.cell(method_row - 1, col)._style)
notes.auto_filter.ref = f"A1:E{notes.max_row}"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT)

reopened = load_workbook(OUTPUT, read_only=True, data_only=False)
check = reopened["Marketing campaign"]
assert check.max_row == 2241 and check.max_column == 26
saved_rows = check.iter_rows(values_only=True)
assert list(next(saved_rows)) == headers
saved_labels = [row[25] for row in saved_rows]
assert saved_labels.count(None) == blank_count
assert len(saved_labels) - saved_labels.count(None) == 2240 - blank_count
assert len(reopened.sheetnames) == len(wb.sheetnames) == 7
print("original labels:", dict(original_labels))
print("rows:", len(rows), "columns:", len(headers), "blank target labels:", blank_count)
print("sheets:", reopened.sheetnames)
print("output:", OUTPUT)
